from __future__ import annotations

import json
import os
import re
import sqlite3
import threading
import uuid
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime
from pathlib import Path
from queue import Queue
from typing import Callable, Iterable

from openpyxl import Workbook, load_workbook


INVALID_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|]')
RESERVED_FILENAMES = {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}


class ExportCancelled(Exception):
    """Raised when the user requests a cooperative task cancellation."""


def normalize_field(value: object) -> str:
    text = "" if value is None else str(value)
    return re.sub(r"\s+", "", text).strip().lower()


def safe_filename(value: object, fallback: str) -> str:
    text = str(value or "").strip()
    text = INVALID_FILENAME_CHARS.sub("_", text).rstrip(". ")
    if not text:
        text = fallback
    if text.upper() in RESERVED_FILENAMES:
        text = f"_{text}"
    return text[:120]


def json_value(value: object) -> object:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def find_header_row(ws, search_limit: int = 30) -> int:
    for row_number, row in enumerate(ws.iter_rows(min_row=1, max_row=search_limit, values_only=True), start=1):
        values = [str(v).strip() for v in row if v is not None and str(v).strip()]
        if len(values) >= 2:
            return row_number
    return 1


def scan_headers(paths: Iterable[str | Path]) -> list[str]:
    names: list[str] = []
    seen: set[str] = set()
    for path in paths:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb[wb.sheetnames[0]]
            header_row = find_header_row(ws)
            headers = [str(v).strip() if v is not None and str(v).strip() else f"未命名列{index + 1}" for index, v in enumerate(next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True)))]
            for header in headers:
                key = normalize_field(header)
                if key and key not in seen:
                    seen.add(key)
                    names.append(header)
        finally:
            wb.close()
    return names


class ImportStore:
    def __init__(self, database_path: str | Path):
        self.database_path = str(database_path)
        # The Tkinter UI performs imports/exports in worker threads. Allow the
        # connection to be shared and serialize access so SQLite's default
        # same-thread guard does not reject those operations.
        self._lock = threading.RLock()
        self.connection = sqlite3.connect(self.database_path, check_same_thread=False)
        self.connection.execute("PRAGMA busy_timeout=30000")
        self.connection.execute("PRAGMA journal_mode=WAL")
        self.connection.execute("PRAGMA synchronous=NORMAL")
        self.connection.execute("PRAGMA temp_store=FILE")
        self.connection.execute("PRAGMA mmap_size=268435456")
        self.connection.executescript(
            """
            CREATE TABLE IF NOT EXISTS import_task (
                task_id TEXT PRIMARY KEY,
                created_at TEXT NOT NULL,
                status TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS source_file (
                file_id INTEGER PRIMARY KEY AUTOINCREMENT,
                task_id TEXT NOT NULL,
                file_name TEXT NOT NULL,
                file_path TEXT NOT NULL,
                sheet_name TEXT NOT NULL,
                header_row INTEGER NOT NULL
            );
            CREATE TABLE IF NOT EXISTS source_column (
                task_id TEXT NOT NULL,
                display_name TEXT NOT NULL,
                normalized_name TEXT NOT NULL,
                PRIMARY KEY (task_id, normalized_name)
            );
            CREATE TABLE IF NOT EXISTS staging_row (
                row_id INTEGER PRIMARY KEY AUTOINCREMENT,
                task_id TEXT NOT NULL,
                file_id INTEGER NOT NULL,
                source_row INTEGER NOT NULL,
                raw_json TEXT NOT NULL
            );
            CREATE INDEX IF NOT EXISTS idx_staging_task ON staging_row(task_id);
            CREATE INDEX IF NOT EXISTS idx_staging_task_row ON staging_row(task_id, row_id);
            """
        )
        self.connection.commit()

    def close(self) -> None:
        with self._lock:
            self.connection.close()

    def import_files(
        self,
        paths: list[str | Path],
        progress: Callable[[str], None] | None = None,
        cancel_event: threading.Event | None = None,
        max_workers: int | None = None,
    ) -> tuple[str, list[str]]:
        with self._lock:
            task_id = uuid.uuid4().hex
            self.connection.execute("INSERT INTO import_task VALUES (?, datetime('now'), 'IMPORTING')", (task_id,))
            self.connection.commit()
            all_headers: dict[str, str] = {}
            imported_rows = 0
            if not paths:
                self.connection.execute("UPDATE import_task SET status='READY' WHERE task_id=?", (task_id,))
                self.connection.commit()
                return task_id, []
            worker_count = max_workers or min(4, len(paths), max(1, os.cpu_count() or 2))
            worker_count = max(1, min(worker_count, len(paths)))
            batch_queue: Queue = Queue(maxsize=worker_count * 4)
            completed_files = 0
            file_ids: dict[int, int] = {}

            def read_file(file_index: int, path_value: str | Path):
                path = Path(path_value)

                def put(message):
                    while True:
                        if cancel_event and cancel_event.is_set():
                            raise ExportCancelled("用户已停止导入任务")
                        try:
                            batch_queue.put(message, timeout=0.5)
                            return
                        except Exception:
                            continue

                wb = None
                try:
                    wb = load_workbook(path, read_only=True, data_only=True)
                    ws = wb[wb.sheetnames[0]]
                    header_row = find_header_row(ws)
                    header_values = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
                    headers: list[str] = []
                    display_headers: list[str] = []
                    for index, value in enumerate(header_values):
                        display = str(value).strip() if value is not None and str(value).strip() else f"未命名列{index + 1}"
                        normalized = normalize_field(display) or f"未命名列{index + 1}"
                        if normalized in headers:
                            normalized = f"{normalized}_{index + 1}"
                        headers.append(normalized)
                        display_headers.append(display)
                    put(("file", file_index, path.name, str(path), ws.title, header_row, headers, display_headers))
                    row_batch: list[tuple[int, str]] = []
                    for row_number, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
                        if cancel_event and cancel_event.is_set():
                            raise ExportCancelled("用户已停止导入任务")
                        if not any(value is not None and str(value).strip() for value in row):
                            continue
                        data = {headers[index]: json_value(row[index] if index < len(row) else None) for index in range(len(headers))}
                        row_batch.append((row_number, json.dumps(data, ensure_ascii=False)))
                        if len(row_batch) >= 5000:
                            put(("batch", file_index, row_batch))
                            row_batch = []
                    if row_batch:
                        put(("batch", file_index, row_batch))
                    put(("done", file_index))
                except Exception as exc:
                    try:
                        batch_queue.put(("error", file_index, exc), timeout=1)
                    except Exception:
                        pass
                finally:
                    if wb is not None:
                        wb.close()

            try:
                with self.connection:
                    if progress:
                        progress(f"并行读取 {len(paths)} 个文件（{worker_count} 个读取线程），SQLite 由单写入通道批量写入...")
                    with ThreadPoolExecutor(max_workers=worker_count, thread_name_prefix="xlsx-reader") as executor:
                        futures = [executor.submit(read_file, index, path_value) for index, path_value in enumerate(paths, start=1)]
                        while completed_files < len(paths):
                            if cancel_event and cancel_event.is_set():
                                raise ExportCancelled("用户已停止导入任务")
                            message = batch_queue.get()
                            kind = message[0]
                            if kind == "file":
                                _, file_index, file_name, file_path, sheet_name, header_row, headers, display_headers = message
                                file_id = self.connection.execute(
                                    "INSERT INTO source_file(task_id, file_name, file_path, sheet_name, header_row) VALUES (?, ?, ?, ?, ?)",
                                    (task_id, file_name, file_path, sheet_name, header_row),
                                ).lastrowid
                                file_ids[file_index] = file_id
                                for normalized, display in zip(headers, display_headers):
                                    all_headers.setdefault(normalized, display)
                            elif kind == "batch":
                                _, file_index, rows = message
                                file_id = file_ids[file_index]
                                values = [(task_id, file_id, source_row, raw_json) for source_row, raw_json in rows]
                                self.connection.executemany(
                                    "INSERT INTO staging_row(task_id, file_id, source_row, raw_json) VALUES (?, ?, ?, ?)", values
                                )
                                imported_rows += len(values)
                                if progress and imported_rows // 100000 != (imported_rows - len(values)) // 100000:
                                    progress(f"已导入 {imported_rows:,} 行...")
                            elif kind == "done":
                                completed_files += 1
                                if progress:
                                    progress(f"已读取完成 {completed_files}/{len(paths)} 个文件...")
                            elif kind == "error":
                                if cancel_event:
                                    cancel_event.set()
                                raise message[2]
                        for future in futures:
                            future.result()
                    self.connection.executemany(
                        "INSERT INTO source_column(task_id, display_name, normalized_name) VALUES (?, ?, ?)",
                        [(task_id, display, normalized) for normalized, display in all_headers.items()],
                    )
                    self.connection.execute("UPDATE import_task SET status='READY' WHERE task_id=?", (task_id,))
                return task_id, [all_headers[key] for key in all_headers]
            except Exception:
                self.connection.execute("DELETE FROM staging_row WHERE task_id=?", (task_id,))
                self.connection.execute("DELETE FROM source_file WHERE task_id=?", (task_id,))
                self.connection.execute("DELETE FROM source_column WHERE task_id=?", (task_id,))
                self.connection.execute("UPDATE import_task SET status='FAILED' WHERE task_id=?", (task_id,))
                self.connection.commit()
                raise

    def iter_rows(self, task_id: str):
        with self._lock:
            cursor = self.connection.execute("SELECT raw_json, source_file.file_name, source_row FROM staging_row JOIN source_file ON source_file.file_id=staging_row.file_id WHERE staging_row.task_id=? ORDER BY row_id", (task_id,))
            for raw_json, file_name, source_row in cursor:
                yield json.loads(raw_json), file_name, source_row

    def iter_rows_with_ids(self, task_id: str):
        with self._lock:
            cursor = self.connection.execute("SELECT row_id, raw_json FROM staging_row WHERE task_id=? ORDER BY row_id", (task_id,))
            for row_id, raw_json in cursor:
                yield row_id, json.loads(raw_json)

    def distinct_date_values(self, task_id: str, year_field: str, month_field: str, day_field: str):
        paths = tuple(json_path(normalize_field(field)) for field in (year_field, month_field, day_field))
        with self._lock:
            cursor = self.connection.execute(
                "SELECT DISTINCT json_extract(raw_json, ?), json_extract(raw_json, ?), json_extract(raw_json, ?) FROM staging_row WHERE task_id=?",
                (*paths, task_id),
            )
            yield from cursor


def load_region_rules(path: str | Path) -> list[dict]:
    config_path = Path(path)
    if not config_path.exists():
        return []
    payload = json.loads(config_path.read_text(encoding="utf-8"))
    return payload.get("regions", [])


def load_customer_names(path: str | Path) -> list[str]:
    """Load preset customer names from a JSON object or a plain JSON list."""
    config_path = Path(path)
    if not config_path.exists():
        return []
    payload = json.loads(config_path.read_text(encoding="utf-8"))
    values = payload.get("customers", []) if isinstance(payload, dict) else payload
    names: list[str] = []
    seen: set[str] = set()
    for value in values or []:
        if isinstance(value, dict):
            value = value.get("name")
        name = str(value or "").strip()
        key = normalize_field(name)
        if name and key not in seen:
            seen.add(key)
            names.append(name)
    return names


def resolve_region(value: object, rules: list[dict]) -> str:
    normalized = normalize_field(value)
    for rule in rules:
        for match in rule.get("matchValues", []):
            if normalize_field(match) == normalized:
                return str(rule.get("name") or rule.get("code") or "未分区")
    return "未分区"


def parse_date_folder(value: object, fallback: str) -> str:
    if value is None or not str(value).strip():
        return fallback
    text = str(value).strip()[:10].replace("/", "-").replace(".", "-")
    if re.match(r"^\d{4}$", text):
        return text
    match = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})$", text)
    if match:
        return f"{int(match.group(1)):04d}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"
    return fallback


def _date_part(value: object) -> int | None:
    """Convert a year/month/day cell to an integer without accepting junk."""
    if value is None or not str(value).strip():
        return None
    try:
        number = int(float(str(value).strip()))
    except (TypeError, ValueError):
        return None
    return number


def json_path(key: str) -> str:
    escaped = key.replace("\\", "\\\\").replace('"', '\\"')
    return f'$."{escaped}"'


def parse_date_parts(year: object, month: object, day: object) -> date | None:
    year_value = _date_part(year)
    month_value = _date_part(month)
    day_value = _date_part(day)
    if year_value is None or month_value is None or day_value is None:
        return None
    try:
        return date(year_value, month_value, day_value)
    except ValueError:
        return None


def available_dates(store: ImportStore, task_id: str, year_field: str | None, month_field: str | None, day_field: str | None) -> list[str]:
    """Return valid distinct dates from the imported task in ISO order."""
    if not year_field or not month_field or not day_field:
        return []
    values: set[date] = set()
    for year, month, day in store.distinct_date_values(task_id, year_field, month_field, day_field):
        parsed = parse_date_parts(year, month, day)
        if parsed:
            values.add(parsed)
    return [item.isoformat() for item in sorted(values)]


def export_task(
    store: ImportStore,
    task_id: str,
    output_root: str | Path,
    selected_fields: list[str],
    customer_field: str,
    region_field: str | None,
    year_field: str | None,
    month_field: str | None = None,
    day_field: str | None = None,
    region_rules: list[dict] | None = None,
    selected_date: str | None = None,
    progress: Callable[[str], None] | None = None,
    date_end: str | None = None,
    customer_names: list[str] | None = None,
    cancel_event: threading.Event | None = None,
) -> dict:
    # Keep the previous API usable for callers that supplied one date field:
    # export_task(..., date_field, region_rules, progress).
    if isinstance(month_field, list) and region_rules is None:
        if callable(day_field) and progress is None:
            progress = day_field
        region_rules = month_field
        month_field = None
        day_field = None
    region_rules = region_rules or []
    default_date = datetime.now().strftime("%Y-%m-%d")
    normalized_selected = [(field, normalize_field(field)) for field in selected_fields]
    customer_key = normalize_field(customer_field)
    customer_filter = {normalize_field(name) for name in (customer_names or []) if normalize_field(name)}
    region_key = normalize_field(region_field) if region_field else None
    year_key = normalize_field(year_field) if year_field else None
    month_key = normalize_field(month_field) if month_field else None
    day_key = normalize_field(day_field) if day_field else None
    selected_date_value = None
    selected_end_value = None
    if selected_date:
        try:
            selected_date_value = date.fromisoformat(selected_date)
        except ValueError:
            raise ValueError(f"无效的导出日期：{selected_date}") from None
    if date_end:
        try:
            selected_end_value = date.fromisoformat(date_end)
        except ValueError:
            raise ValueError(f"无效的结束日期：{date_end}") from None
        if selected_date_value and selected_end_value < selected_date_value:
            selected_date_value, selected_end_value = selected_end_value, selected_date_value
    range_folder = None
    if selected_date_value:
        range_folder = selected_date_value.isoformat()
        if selected_end_value and selected_end_value != selected_date_value:
            range_folder = f"{range_folder}至{selected_end_value.isoformat()}"
    root = Path(output_root)
    root.mkdir(parents=True, exist_ok=True)
    row_count = 0
    group_count = 0
    output_count = 0
    bucket_batch: list[tuple[int, str, str, str]] = []
    current_key: tuple[str, str, str] | None = None
    workbook = None
    worksheet = None
    current_rows = 0
    part_number = 0

    def close_workbook(key: tuple[str, str, str] | None, current_workbook, current_part: int):
        nonlocal output_count
        if not current_workbook or not key:
            return
        folder_date, region, customer = key
        target_dir = root / safe_filename(folder_date, default_date) / safe_filename(region, "未分区")
        target_dir.mkdir(parents=True, exist_ok=True)
        suffix = "" if current_part == 1 else f"_part{current_part}"
        target_file = target_dir / f"{safe_filename(customer, '未识别客户')}{suffix}.xlsx"
        current_workbook.save(target_file)
        current_workbook.close()
        output_count += 1
        if progress:
            progress(f"已生成 {target_file}")

    def check_cancelled():
        if cancel_event and cancel_event.is_set():
            if workbook:
                try:
                    workbook.close()
                except Exception:
                    pass
            try:
                connection.execute("DROP TABLE IF EXISTS temp.export_bucket")
                connection.commit()
            except Exception:
                pass
            raise ExportCancelled("用户已停止导出任务")

    with store._lock:
        connection = store.connection
        connection.execute("DROP TABLE IF EXISTS temp.export_bucket")
        connection.execute("CREATE TEMP TABLE export_bucket (row_id INTEGER PRIMARY KEY, folder_date TEXT NOT NULL, region TEXT NOT NULL, customer TEXT NOT NULL)")
        for row_id, row in store.iter_rows_with_ids(task_id):
            check_cancelled()
            customer = str(row.get(customer_key) or "").strip() or "未识别客户"
            if customer_filter and normalize_field(customer) not in customer_filter:
                continue
            region = resolve_region(row.get(region_key), region_rules) if region_key else "未分区"
            parsed_date = parse_date_parts(row.get(year_key), row.get(month_key), row.get(day_key)) if year_key and month_key and day_key else None
            if selected_date_value and (parsed_date is None or parsed_date < selected_date_value or (selected_end_value and parsed_date > selected_end_value)):
                continue
            # A selected date range is one filter/group. Keep all matching
            # days for a customer in the same workbook instead of creating
            # one date folder per day.
            folder_date = range_folder or (parsed_date.isoformat() if parsed_date else default_date)
            bucket_batch.append((row_id, folder_date, region, customer))
            row_count += 1
            if len(bucket_batch) >= 10000:
                connection.executemany("INSERT INTO temp.export_bucket(row_id, folder_date, region, customer) VALUES (?, ?, ?, ?)", bucket_batch)
                bucket_batch.clear()
                if progress and row_count % 100000 == 0:
                    progress(f"已准备 {row_count:,} 行导出数据...")
        if bucket_batch:
            connection.executemany("INSERT INTO temp.export_bucket(row_id, folder_date, region, customer) VALUES (?, ?, ?, ?)", bucket_batch)
        connection.execute("CREATE INDEX temp.idx_export_bucket_group ON export_bucket(folder_date, region, customer, row_id)")
        connection.commit()

        cursor = connection.execute(
            "SELECT export_bucket.folder_date, export_bucket.region, export_bucket.customer, staging_row.raw_json "
            "FROM temp.export_bucket JOIN staging_row ON staging_row.row_id=export_bucket.row_id "
            "ORDER BY export_bucket.folder_date, export_bucket.region, export_bucket.customer, export_bucket.row_id"
        )
        for folder_date, region, customer, raw_json in cursor:
            check_cancelled()
            key = (folder_date, region, customer)
            if key != current_key:
                close_workbook(current_key, workbook, part_number)
                current_key = key
                group_count += 1
                part_number = 1
                workbook = Workbook(write_only=True)
                worksheet = workbook.create_sheet("数据")
                worksheet.freeze_panes = "A2"
                worksheet.append(selected_fields)
                current_rows = 0
            if current_rows >= 1_048_575:
                close_workbook(current_key, workbook, part_number)
                part_number += 1
                workbook = Workbook(write_only=True)
                worksheet = workbook.create_sheet("数据")
                worksheet.freeze_panes = "A2"
                worksheet.append(selected_fields)
                current_rows = 0
            row = json.loads(raw_json)
            worksheet.append([row.get(key_name) for _, key_name in normalized_selected])
            current_rows += 1
        close_workbook(current_key, workbook, part_number)
        connection.execute("DROP TABLE IF EXISTS temp.export_bucket")
        connection.commit()
    return {"rows": row_count, "files": output_count, "groups": group_count, "errors": []}
